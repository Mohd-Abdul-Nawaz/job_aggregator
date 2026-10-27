"""
resume_parser.py
-----------------
Extracts raw text from an uploaded resume file (PDF or DOCX).

Design goal: this module is 100% generic. It does NOT look for specific
skills, keywords, or sections. It just pulls out plain text so that the
matching_engine module can compare it against any job description,
for any field/industry.
"""

import io
import pdfplumber
import docx  # python-docx


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract all text from a PDF resume using pdfplumber."""
    text_chunks = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text)
    return "\n".join(text_chunks)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract all text from a DOCX resume using python-docx."""
    document = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # Also pull text out of any tables (many resumes use table layouts)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs)


def parse_resume(uploaded_file) -> str:
    """
    Generic entry point. Accepts a Streamlit UploadedFile object,
    detects the file type by extension, and returns extracted plain text.

    Raises a ValueError for unsupported file types so the UI can show
    a friendly error instead of crashing.
    """
    if uploaded_file is None:
        return ""

    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()

    if filename.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif filename.endswith(".docx"):
        text = extract_text_from_docx(file_bytes)
    else:
        raise ValueError(
            f"Unsupported resume format: '{filename}'. Please upload a PDF or DOCX file."
        )

    # Normalize whitespace a bit for cleaner downstream matching
    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    return cleaned
